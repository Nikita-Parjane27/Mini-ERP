import csv

import pytest
from openpyxl import load_workbook

from mini_erp import inventory as inv, reports
from mini_erp.errors import ERPError

from conftest import count, fill_stock


def write(path, text):
    path.write_text(text, encoding="utf-8")
    return str(path)


def test_import_loads_good_rows_and_rejects_bad_ones(conn, tmp_path):
    f = write(tmp_path / "items.csv",
              "sku,name,unit,item_type,reorder_level,unit_cost\n"
              "A1,Item A,kg,RAW,5,10\n"
              ",No sku,kg,RAW,1,1\n"
              "A2,,kg,RAW,1,1\n"
              "A3,Bad type,kg,CHEMICAL,1,1\n"
              "A4,Neg reorder,kg,RAW,-5,1\n"
              "A5,Bad cost,kg,RAW,1,abc\n"
              "A6,Good,pcs,FINISHED,2,0\n")
    before = count(conn, "items")
    res = inv.import_items_csv(conn, f, str(tmp_path / "errors.csv"))
    assert res.added == 2 and len(res.errors) == 5
    assert count(conn, "items") == before + 2


def test_import_catches_duplicates_in_file_and_in_database(conn, tmp_path):
    f = write(tmp_path / "items.csv", "sku,name\nB1,First\nB1,Second\nCU,Already in DB\n")
    res = inv.import_items_csv(conn, f)
    assert res.added == 1
    reasons = " ".join(msg for _, _, msg in res.errors)
    assert "Duplicate SKU 'B1'" in reasons and "already exists" in reasons


def test_import_writes_error_report_csv(conn, tmp_path):
    f = write(tmp_path / "items.csv", "sku,name\nC1,\n")
    err = tmp_path / "errors.csv"
    inv.import_items_csv(conn, f, str(err))
    rows = list(csv.DictReader(open(err, encoding="utf-8")))
    assert rows[0]["line"] == "2" and "name" in rows[0]["error"].lower()


def test_import_fails_fast_when_required_column_missing(conn, tmp_path):
    f = write(tmp_path / "items.csv", "sku,unit\nD1,kg\n")
    with pytest.raises(ERPError, match="missing required column"):
        inv.import_items_csv(conn, f)


def test_low_stock_report_lists_items_at_or_below_reorder_level(conn):
    inv.adjust_stock(conn, "CU", 50)    # equals reorder level -> low
    inv.adjust_stock(conn, "CORE", 21)  # above reorder level -> ok
    assert [r["sku"] for r in reports.low_stock_rows(conn)] == ["CT", "CU"]  # sorted by biggest shortfall first


def test_valuation_uses_last_purchase_rate(conn):
    fill_stock(conn, cu=10, core=10)  # 10 x 700 and 10 x 90
    values = {r["sku"]: r["value"] for r in reports.valuation_rows(conn)}
    assert values["CU"] == 7000 and values["CORE"] == 900


def test_purchase_summary_counts_only_received_pos(conn):
    fill_stock(conn, cu=10, core=10)                            # received: 7000 + 900
    inv.create_po(conn, "Shree Metals", [("CU", 999, 700)])      # still OPEN, ignored
    row = reports.purchase_summary_rows(conn)[0]
    assert row["supplier"] == "Shree Metals" and row["po_count"] == 1 and row["total_value"] == 7900


def test_excel_report_has_all_sheets(conn, tmp_path):
    fill_stock(conn)
    path = reports.export_excel(conn, str(tmp_path / "r.xlsx"))
    wb = load_workbook(path)
    assert wb.sheetnames == ["Stock", "Low Stock", "Valuation", "Purchase Summary"]
    assert wb["Stock"]["A1"].value == "sku"


def test_daily_job_saves_report_and_returns_low_items(conn, tmp_path):
    inv.adjust_stock(conn, "CU", 5)
    path, low = reports.daily_job(conn, str(tmp_path / "out"))
    assert path.endswith(".xlsx") and any(r["sku"] == "CU" for r in low)
