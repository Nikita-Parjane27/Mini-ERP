"""Read-only queries, plus CSV / Excel / JSON exports."""
import csv
import json
import logging
import os
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

log = logging.getLogger("mini_erp")


def stock_rows(conn):
    return conn.execute(
        "SELECT sku, name, item_type, unit, stock_qty, reorder_level, "
        "CASE WHEN stock_qty <= reorder_level THEN 'LOW' ELSE 'OK' END AS status "
        "FROM items ORDER BY item_type DESC, sku"
    ).fetchall()


def get_stock(conn, sku):
    return conn.execute(
        "SELECT sku, name, item_type, unit, stock_qty, reorder_level, "
        "CASE WHEN stock_qty <= reorder_level THEN 'LOW' ELSE 'OK' END AS status "
        "FROM items WHERE sku = ?",
        (sku.strip().upper(),),
    ).fetchone()


def low_stock_rows(conn):
    return conn.execute(
        "SELECT sku, name, unit, stock_qty, reorder_level, (reorder_level - stock_qty) AS shortfall "
        "FROM items WHERE stock_qty <= reorder_level ORDER BY shortfall DESC, sku"
    ).fetchall()


def valuation_rows(conn):
    """Stock value = quantity x last purchase rate."""
    return conn.execute(
        "SELECT sku, name, stock_qty, unit_cost, ROUND(stock_qty * unit_cost, 2) AS value "
        "FROM items ORDER BY value DESC, sku"
    ).fetchall()


def purchase_summary_rows(conn):
    """JOIN + GROUP BY: received purchases per supplier."""
    return conn.execute(
        "SELECT s.name AS supplier, COUNT(DISTINCT po.id) AS po_count, "
        "ROUND(SUM(l.qty * l.rate), 2) AS total_value "
        "FROM purchase_orders po "
        "JOIN suppliers s ON s.id = po.supplier_id "
        "JOIN po_lines l ON l.po_id = po.id "
        "WHERE po.status = 'RECEIVED' "
        "GROUP BY s.id ORDER BY total_value DESC"
    ).fetchall()


def po_rows(conn):
    return conn.execute(
        "SELECT po.po_no, s.name AS supplier, po.status, po.created_on, po.received_on, "
        "ROUND(SUM(l.qty * l.rate), 2) AS total "
        "FROM purchase_orders po JOIN suppliers s ON s.id = po.supplier_id "
        "JOIN po_lines l ON l.po_id = po.id GROUP BY po.id ORDER BY po.id"
    ).fetchall()


def ledger_rows(conn, sku=None, limit=None):
    sql = (
        "SELECT l.id, l.created_at, i.sku, l.movement_type, l.qty_change, l.note "
        "FROM stock_ledger l JOIN items i ON i.id = l.item_id "
    )
    params = []
    if sku:
        sql += "WHERE i.sku = ? "
        params.append(sku.strip().upper())
    sql += "ORDER BY l.id"
    if limit:
        sql = f"SELECT * FROM ({sql} DESC LIMIT ?) ORDER BY id"
        params.append(int(limit))
    return conn.execute(sql, params).fetchall()


def verify_ledger(conn):
    """Audit: recompute stock from the ledger and list items where it does not match."""
    return conn.execute(
        "SELECT i.sku, i.stock_qty AS recorded, ROUND(COALESCE(SUM(l.qty_change), 0), 6) AS from_ledger "
        "FROM items i LEFT JOIN stock_ledger l ON l.item_id = i.id "
        "GROUP BY i.id "
        "HAVING ABS(i.stock_qty - COALESCE(SUM(l.qty_change), 0)) > 0.000001"
    ).fetchall()


# ---------- exports ----------

def _rows_to_lists(rows):
    if not rows:
        return []
    return [list(rows[0].keys())] + [list(r) for r in rows]


def write_csv(rows, path):
    data = _rows_to_lists(rows)
    if not data:
        return 0
    with open(path, "w", newline="", encoding="utf-8") as f:
        csv.writer(f).writerows(data)
    return len(rows)


def export_excel(conn, path):
    """One workbook: Stock, Low Stock, Valuation, Purchase Summary."""
    sheets = {
        "Stock": stock_rows(conn),
        "Low Stock": low_stock_rows(conn),
        "Valuation": valuation_rows(conn),
        "Purchase Summary": purchase_summary_rows(conn),
    }
    wb = Workbook()
    wb.remove(wb.active)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for title, rows in sheets.items():
        ws = wb.create_sheet(title)
        data = _rows_to_lists(rows) or [["(no data)"]]
        for r in data:
            ws.append(r)
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = header_fill
        for col_idx in range(1, ws.max_column + 1):
            letter = get_column_letter(col_idx)
            width = max(len(str(c.value)) if c.value is not None else 0 for c in ws[letter])
            ws.column_dimensions[letter].width = min(width + 3, 45)
        ws.freeze_panes = "A2"
    if sheets["Valuation"]:
        ws = wb["Valuation"]
        last = ws.max_row + 1
        ws.cell(row=last, column=1, value="TOTAL").font = Font(bold=True)
        total = ws.cell(row=last, column=5, value=f"=SUM(E2:E{last - 1})")
        total.font = Font(bold=True)
    wb.save(path)
    return path


def export_json(conn, path="stock.json"):
    data = {
        "generated_on": date.today().isoformat(),
        "stock": [dict(r) for r in stock_rows(conn)],
        "low_stock": [dict(r) for r in low_stock_rows(conn)],
        "valuation": [dict(r) for r in valuation_rows(conn)],
        "purchase_summary": [dict(r) for r in purchase_summary_rows(conn)],
    }
    with open(path, "w", encoding="utf-8") as f:
        json.dump(data, f, indent=2)
    return path


def daily_job(conn, out_dir="reports"):
    """Run once a day (Task Scheduler / cron): save the Excel report and log low-stock items."""
    os.makedirs(out_dir, exist_ok=True)
    path = export_excel(conn, os.path.join(out_dir, f"daily_report_{date.today().isoformat()}.xlsx"))
    low = low_stock_rows(conn)
    if low:
        names = ", ".join(f"{r['sku']} ({r['stock_qty']:g}/{r['reorder_level']:g})" for r in low)
        log.warning("DAILY JOB: %d item(s) at or below reorder level: %s", len(low), names)
    else:
        log.info("DAILY JOB: all items above reorder level")
    return path, low
